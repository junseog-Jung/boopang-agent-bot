import csv
import io
import json
from decimal import Decimal

import pytest
from openpyxl import Workbook

from boopang.config import Config
from boopang.exporter import to_csv, to_dict, to_json, to_xlsx, safe_cell
from boopang.inspector import normalize_phone, parse_discount, parse_int, parse_money
from boopang.models import ParseError, Status
from boopang.pipeline import process
from boopang.sample import sample_xlsx

CFG = Config()
HEAD = "주문번호,SKU,품목명,주문수량,검수수량,단가,할인율,연락처,출고일,운송장번호,배송상태\n"


def run(body: str, cfg=CFG, head=HEAD):
    return process((head + body).encode("utf-8"), "t.csv", cfg)


def codes(res):
    return {(i.row_no, i.code) for i in res.issues}


# ---- 파서 ----
def test_header_aliases_english_and_korean():
    res = process(b"OrderNo,Item Code,Qty,Price\nX1,S1,2,1000\n", "a.csv", CFG)
    assert res.lines[0].order_id == "X1" and res.lines[0].qty_ordered == 2


def test_missing_required_header():
    with pytest.raises(ParseError):
        process("주문번호,품목명\nA,B\n".encode(), "a.csv", CFG)


def test_cp949_csv():
    raw = (HEAD + "A1,S1,품목,1,1,1000,0,,,,\n").encode("cp949")
    assert process(raw, "a.csv", CFG).lines[0].name == "품목"


def test_unsupported_extension_and_bad_xlsx():
    with pytest.raises(ParseError):
        process(b"x", "a.exe", CFG)
    with pytest.raises(ParseError):
        process(b"not a zip", "a.xlsx", CFG)


def test_row_limit():
    body = "".join(f"A{i},S,n,1,1,100,0,,,,\n" for i in range(10))
    with pytest.raises(ParseError):
        run(body, Config(max_rows=5))


def test_blank_rows_skipped_and_empty_file():
    assert len(run("A1,S1,n,1,1,100,0,,,,\n,,,,,,,,,,\n").lines) == 1
    with pytest.raises(ParseError):
        run("")


# ---- 값 변환 ----
@pytest.mark.parametrize("v,exp", [("10%", "0.1"), (10, "0.1"), (0.1, "0.1"), ("", "0"), (None, "0")])
def test_discount(v, exp):
    d, ok = parse_discount(v)
    assert ok and d == Decimal(exp)


@pytest.mark.parametrize("v", ["abc", "150%", -5, "nan"])
def test_discount_invalid(v):
    assert parse_discount(v)[1] is False


def test_money_and_int_edge_cases():
    assert parse_money("12,000원") == (Decimal("12000"), True)
    assert parse_money("NaN")[1] is False and parse_money("Infinity")[1] is False
    assert parse_money("9" * 20)[1] is False
    assert parse_int("1.5")[1] is False and parse_int(True)[1] is False
    assert parse_int("1,000") == (1000, True) and parse_int(3.0) == (3, True)


def test_phone_normalization():
    assert normalize_phone("01012345678") == ("010-1234-5678", True)
    assert normalize_phone("+82 10-1234-5678") == ("010-1234-5678", True)
    assert normalize_phone("021234567") == ("02-123-4567", True)
    assert normalize_phone("1234")[1] is False


# ---- 검수 ----
def test_inspection_issues():
    res = run(
        "A1,S1,n,0,0,100,0,,,,\n"      # 수량 0 → ERROR
        "A2,S1,n,1,1,-5,0,,,,\n"       # 음수 단가 → ERROR
        ",S1,n,1,1,100,0,,,,\n"        # 주문번호 없음
        "A3,,n,1,1,100,0,,,,\n"        # SKU 없음
        "A4,S1,n,1,1,100,0,,2026-13-45,,\n"  # 날짜 오류 → WARN
        "A5,S1,n,2,5,100,0,,,,\n"      # 초과
        "A6,S1,n,5,2,100,0,,,,\n"      # 부족
        "A7,S1,n,1,1,0,0,,,,\n"        # 0원
    )
    c = codes(res)
    assert (2, "INVALID_QTY") in c and (3, "INVALID_PRICE") in c
    assert (4, "MISSING_ORDER_ID") in c and (5, "MISSING_SKU") in c
    assert (6, "INVALID_DATE") in c and (7, "OVERAGE") in c
    assert (8, "SHORTAGE") in c and (9, "ZERO_PRICE") in c
    assert [l.status for l in res.lines[:4]] == [Status.HOLD] * 4


def test_duplicate_line_warns():
    res = run("A1,S1,n,1,1,100,0,,,,\nA1,S1,n,1,1,100,0,,,,\n")
    assert (3, "DUPLICATE_LINE") in codes(res) and res.lines[1].line_no == 2


# ---- 금액 ----
def test_pricing_vat_excluded_with_discount_and_rounding():
    res = run("A1,S1,n,3,3,3333,10%,,,,\n", Config(free_shipping_threshold=Decimal("0")))
    l = res.lines[0]
    assert l.supply_amt == Decimal("8999")     # 3*3333*0.9 = 8999.1 → 8999
    assert l.vat_amt == Decimal("900")          # 899.9 → 900
    assert l.total_amt == Decimal("9899")


def test_pricing_half_up_not_bankers():
    res = run("A1,S1,n,1,1,2.5,0,,,,\n", Config(free_shipping_threshold=Decimal("0")))
    assert res.lines[0].supply_amt == Decimal("3")


def test_pricing_vat_included():
    res = run("A1,S1,n,1,1,11000,0,,,,\n", Config(price_includes_vat=True, free_shipping_threshold=Decimal("0")))
    l = res.lines[0]
    assert (l.supply_amt, l.vat_amt, l.total_amt) == (Decimal("10000"), Decimal("1000"), Decimal("11000"))


def test_billing_uses_received_qty_capped_by_ordered():
    res = run("A1,S1,n,10,4,1000,0,,,,\nA2,S1,n,2,5,1000,0,,,,\n", Config(free_shipping_threshold=Decimal("0")))
    assert res.lines[0].supply_amt == Decimal("4000")
    assert res.lines[1].supply_amt == Decimal("2000")


def test_shipping_fee_threshold_and_exclusions():
    res = run("A1,S1,n,1,1,10000,0,,,,\nB1,S1,n,10,10,10000,0,,,,\n")
    a, b = res.orders
    assert a.shipping_fee == Decimal("3000") and a.total_amt == Decimal("14000")
    assert b.shipping_fee == Decimal("0")
    res = run("C1,S1,n,1,1,10000,0,,,,취소\n")
    assert res.orders[0].total_amt == 0 and res.orders[0].shipping_fee == 0  # 취소 주문은 배송비 없음


# ---- 상태 ----
@pytest.mark.parametrize("row,expected", [
    ("A,S,n,5,5,100,0,,,,", Status.READY),
    ("A,S,n,5,,100,0,,,,", Status.INSPECT_PENDING),
    ("A,S,n,5,0,100,0,,,,", Status.BACKORDER),
    ("A,S,n,5,3,100,0,,,,", Status.PARTIAL),
    ("A,S,n,5,5,100,0,,,TRK1,", Status.SHIPPING),
    ("A,S,n,5,5,100,0,,,TRK1,배송완료", Status.DELIVERED),
    ("A,S,n,5,5,100,0,,,,취소", Status.CANCELLED),
    ("A,S,n,5,5,100,0,,,,반품", Status.RETURNED),
    ("A,S,n,5,5,100,0,,,TRK1,delivered", Status.DELIVERED),
])
def test_line_status(row, expected):
    assert run(row + "\n").lines[0].status is expected


def test_status_conflict_warnings():
    res = run("A,S,n,1,1,100,0,,,,배송완료\n")
    assert (2, "STATUS_CONFLICT") in codes(res)
    res = run("A,S,n,1,1,100,0,,,,뭔가이상\n")
    assert (2, "UNKNOWN_STATUS") in codes(res)


def test_order_rollup():
    res = run(
        "O1,S1,n,1,1,100,0,,,T1,배송완료\nO1,S2,n,1,1,100,0,,,T2,\n"   # 완료+배송중 → SHIPPING
        "O2,S1,n,1,1,100,0,,,,\nO2,S2,n,1,0,100,0,,,,\n"             # READY+BACKORDER → PARTIAL
        "O3,S1,n,1,1,100,0,,,,\nO3,S2,n,0,1,100,0,,,,\n"             # 오류 포함 → HOLD
        "O4,S1,n,1,1,100,0,,,,취소\nO4,S2,n,1,1,100,0,,,,취소\n"      # 전부 취소
    )
    st = {o.order_id: o.status for o in res.orders}
    assert st == {"O1": Status.SHIPPING, "O2": Status.PARTIAL, "O3": Status.HOLD, "O4": Status.CANCELLED}


# ---- 내보내기 ----
def test_csv_export_columns_and_bom():
    out = to_csv(process(sample_xlsx(), "s.xlsx", CFG))
    assert out.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(out.decode("utf-8-sig"))))
    assert len(rows) == 11 and rows[0]["DISC_RATE"] == "10" and "E" not in rows[0]["DISC_RATE"]


def test_formula_injection_neutralized():
    res = run('A1,S1,"=HYPERLINK(""http://x"")",1,1,100,0,,,,\n')
    csv_text = to_csv(res).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(to_xlsx(res)))
    cell = [c.value for r in wb["WMS라인"].iter_rows(min_row=2) for c in r if c.value and "HYPERLINK" in str(c.value)][0]
    assert cell.startswith("'=")
    # 숫자·정상 문자열은 건드리지 않음
    assert safe_cell("abc") == "abc" and safe_cell(5) == 5


def test_json_and_xlsx_export():
    res = process(sample_xlsx(), "s.xlsx", CFG)
    data = json.loads(to_json(res))
    assert data["schema_version"] == "1.0" and len(data["orders"]) == 9
    assert data["summary"] == to_dict(res)["summary"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(to_xlsx(res)))
    assert wb.sheetnames == ["요약", "주문", "WMS라인", "검수이슈"]


def test_sample_summary_totals():
    s = process(sample_xlsx(), "s.xlsx", CFG).summary()
    assert s["issues"]["ERROR"] == 2 and s["line_count"] == 11 and s["order_count"] == 9


def test_xlsx_numeric_ids_and_dates():
    wb = Workbook(); ws = wb.active
    ws.append(["주문번호", "SKU", "주문수량", "단가", "출고일"])
    from datetime import datetime
    ws.append([1001.0, "S1", 2.0, 1500, datetime(2026, 10, 1)])
    buf = io.BytesIO(); wb.save(buf)
    l = process(buf.getvalue(), "n.xlsx", CFG).lines[0]
    assert l.order_id == "1001" and l.qty_ordered == 2 and l.ship_date.isoformat() == "2026-10-01"


# ---- 보안/견고성 ----
def test_zip_bomb_rejected(monkeypatch):
    import zipfile
    from boopang import parser
    monkeypatch.setattr(parser, "MAX_UNCOMPRESSED_BYTES", 1000)
    with pytest.raises(ParseError, match="너무 큰"):
        process(sample_xlsx(), "s.xlsx", CFG)


def test_corrupt_xlsx_internals_rejected():
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<broken")
        z.writestr("xl/workbook.xml", "<broken")
    with pytest.raises(ParseError):
        process(buf.getvalue(), "s.xlsx", CFG)


def test_huge_csv_cell_rejected():
    big = '주문번호,SKU,주문수량,단가\nA,"' + "x" * 200000 + '",1,1\n'
    with pytest.raises(ParseError):
        process(big.encode(), "a.csv", CFG)
