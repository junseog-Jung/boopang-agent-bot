"""업로드 파일(.xlsx/.csv) → 정규화된 행."""
from __future__ import annotations

import csv
import io
import re
import zipfile

from openpyxl import load_workbook

from .config import Config
from .models import ParseError, RawRow

# 정규 필드명 → 인식 가능한 헤더 별칭 (공백/대소문자/특수문자 제거 후 비교)
ALIASES = {
    "order_id": ["주문번호", "주문no", "주문id", "orderid", "orderno", "ordernumber", "order"],
    "sku": ["sku", "품목코드", "상품코드", "품번", "itemcode", "productcode", "자재코드"],
    "name": ["품목명", "상품명", "품명", "itemname", "productname", "name", "description"],
    "qty_ordered": ["주문수량", "수량", "orderqty", "orderedqty", "qty", "quantity", "ordqty"],
    "qty_received": ["검수수량", "입고수량", "실수량", "실사수량", "receivedqty", "rcvqty", "received", "inspectedqty"],
    "unit_price": ["단가", "가격", "판매가", "unitprice", "price"],
    "discount_rate": ["할인율", "할인", "discount", "discountrate"],
    "recipient": ["수령인", "받는분", "받는사람", "수취인", "recipient", "receiver", "consignee"],
    "phone": ["연락처", "전화번호", "휴대폰", "핸드폰", "phone", "tel", "mobile"],
    "address": ["주소", "배송지", "배송주소", "address", "shippingaddress"],
    "ship_date": ["출고일", "발송일", "배송일", "shipdate", "shippeddate", "shippingdate"],
    "tracking_no": ["운송장번호", "운송장", "송장번호", "trackingno", "trackingnumber", "tracking", "waybill"],
    "raw_status": ["배송상태", "상태", "status", "deliverystatus", "shippingstatus"],
    "memo": ["비고", "메모", "memo", "note", "notes", "remark"],
}
REQUIRED_HEADERS = ["order_id", "sku", "qty_ordered", "unit_price"]

_NORMALIZE = re.compile(r"[\s_\-./()\[\]]+")
_ALIAS_LOOKUP = {}
for _field, _names in ALIASES.items():
    for _n in _names:
        _ALIAS_LOOKUP[_NORMALIZE.sub("", _n).lower()] = _field


def _field_for(header) -> str | None:
    if header is None:
        return None
    return _ALIAS_LOOKUP.get(_NORMALIZE.sub("", str(header)).lower())


def _find_header(rows: list) -> tuple[int, dict]:
    """처음 10행 중 인식되는 헤더가 가장 많은 행을 헤더로 선택."""
    best_idx, best_map = -1, {}
    for idx, row in enumerate(rows[:10]):
        mapping = {}
        for col, cell in enumerate(row):
            f = _field_for(cell)
            if f and f not in mapping.values():
                mapping[col] = f
        if len(mapping) > len(best_map):
            best_idx, best_map = idx, mapping
    missing = [f for f in REQUIRED_HEADERS if f not in best_map.values()]
    if best_idx < 0 or missing:
        labels = {"order_id": "주문번호", "sku": "SKU/품목코드", "qty_ordered": "주문수량", "unit_price": "단가"}
        raise ParseError("필수 컬럼을 찾을 수 없습니다: " + ", ".join(labels[m] for m in missing or REQUIRED_HEADERS))
    return best_idx, best_map


MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024  # xlsx 압축 폭탄 방어


def _read_csv_rows(content: bytes) -> list:
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ParseError("CSV 인코딩을 인식할 수 없습니다 (UTF-8 또는 CP949 필요).")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    try:
        return [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    except csv.Error as exc:
        raise ParseError("CSV 형식이 올바르지 않습니다 (셀이 너무 크거나 따옴표가 깨졌습니다).") from exc


def _read_xlsx_rows(content: bytes, max_rows: int) -> list:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            if sum(i.file_size for i in zf.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ParseError("압축 해제 시 크기가 너무 큰 파일입니다.")
    except zipfile.BadZipFile as exc:
        raise ParseError("올바른 .xlsx 파일이 아닙니다.") from exc
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 은 손상 파일에 다양한 예외를 던진다
        raise ParseError("올바른 .xlsx 파일이 아닙니다.") from exc
    try:
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > max_rows + 20:
                raise ParseError(f"행 수가 제한({max_rows}행)을 초과했습니다.")
        return rows
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("엑셀 시트를 읽는 중 오류가 발생했습니다.") from exc
    finally:
        wb.close()


def parse_file(content: bytes, filename: str, config: Config) -> list[RawRow]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        rows = _read_xlsx_rows(content, config.max_rows)
    elif name.endswith(".csv"):
        rows = _read_csv_rows(content)
    else:
        raise ParseError("지원하지 않는 파일 형식입니다 (.xlsx, .csv 만 가능).")

    if not rows:
        raise ParseError("빈 파일입니다.")
    header_idx, mapping = _find_header(rows)
    out: list[RawRow] = []
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        data = {f: (row[c] if c < len(row) else None) for c, f in mapping.items()}
        if all(v is None or str(v).strip() == "" for v in data.values()):
            continue  # 완전히 빈 행 무시
        out.append(RawRow(row_no=offset, data=data))
        if len(out) > config.max_rows:
            raise ParseError(f"행 수가 제한({config.max_rows}행)을 초과했습니다.")
    if not out:
        raise ParseError("데이터 행이 없습니다.")
    return out
