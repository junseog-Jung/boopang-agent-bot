"""물류 DB(WMS) 연동 포맷 출력 — CSV / JSON / XLSX 리포트."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import STATUS_LABELS, Line, Order, Result, Status

SCHEMA_VERSION = "1.0"
WMS_COLUMNS = [
    "ORDER_NO", "LINE_NO", "SKU", "ITEM_NAME", "ORD_QTY", "RCV_QTY", "UNIT_PRICE", "DISC_RATE",
    "SUPPLY_AMT", "VAT_AMT", "TOTAL_AMT", "STATUS_CD", "RECIPIENT", "TEL", "ADDR", "SHIP_DT",
    "TRACKING_NO", "ISSUE_CNT",
]
_DANGEROUS_PREFIX = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """엑셀/CSV 수식 인젝션 방어: 위험한 접두 문자가 있는 문자열 앞에 ' 를 붙인다."""
    if isinstance(value, str) and value.startswith(_DANGEROUS_PREFIX):
        return "'" + value
    return value


def _money(d: Decimal) -> str:
    return str(d.quantize(Decimal("1")))


def line_record(l: Line) -> dict:
    """WMS 라인 레코드 (문자열 값은 원본 그대로, 정리는 출력 단계에서)."""
    return {
        "ORDER_NO": l.order_id,
        "LINE_NO": l.line_no,
        "SKU": l.sku,
        "ITEM_NAME": l.name,
        "ORD_QTY": l.qty_ordered if l.qty_ordered is not None else "",
        "RCV_QTY": l.qty_received if l.qty_received is not None else "",
        "UNIT_PRICE": _money(l.unit_price) if l.unit_price is not None else "",
        "DISC_RATE": format((l.discount_rate * 100).normalize(), "f") if l.discount_rate else "0",
        "SUPPLY_AMT": _money(l.supply_amt),
        "VAT_AMT": _money(l.vat_amt),
        "TOTAL_AMT": _money(l.total_amt),
        "STATUS_CD": l.status.value,
        "RECIPIENT": l.recipient,
        "TEL": l.phone,
        "ADDR": l.address,
        "SHIP_DT": l.ship_date.isoformat() if l.ship_date else "",
        "TRACKING_NO": l.tracking_no,
        "ISSUE_CNT": l.issue_count,
    }


def order_record(o: Order) -> dict:
    return {
        "ORDER_NO": o.order_id,
        "STATUS_CD": o.status.value,
        "STATUS_LABEL": STATUS_LABELS[o.status],
        "LINE_CNT": len(o.lines),
        "SUPPLY_AMT": _money(o.supply_amt),
        "VAT_AMT": _money(o.vat_amt),
        "SHIPPING_FEE": _money(o.shipping_fee),
        "TOTAL_AMT": _money(o.total_amt),
    }


def to_csv(result: Result) -> bytes:
    """WMS 라인 CSV (UTF-8 BOM — 엑셀 호환)."""
    buf = io.StringIO(newline="")
    w = csv.DictWriter(buf, fieldnames=WMS_COLUMNS, lineterminator="\r\n")
    w.writeheader()
    for l in result.lines:
        rec = {k: safe_cell(v) for k, v in line_record(l).items()}
        w.writerow(rec)
    return buf.getvalue().encode("utf-8-sig")


def to_dict(result: Result) -> dict:
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "summary": result.summary(),
        "orders": [dict(order_record(o), lines=[line_record(l) for l in o.lines]) for o in result.orders],
        "issues": [
            {"row": i.row_no, "severity": i.severity.value, "code": i.code, "message": i.message,
             "order_id": i.order_id, "sku": i.sku}
            for i in result.issues
        ],
    }


def to_json(result: Result) -> bytes:
    return json.dumps(to_dict(result), ensure_ascii=False, indent=2).encode("utf-8")


def _write_sheet(ws, headers: list, rows: list) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    for idx, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[idx - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 40)
    ws.freeze_panes = "A2"


def to_xlsx(result: Result) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    s = result.summary()
    summary_rows = [
        ["원본 파일", s["source"]], ["라인 수", s["line_count"]], ["주문 수", s["order_count"]],
        ["오류(ERROR)", s["issues"]["ERROR"]], ["경고(WARN)", s["issues"]["WARN"]], ["정보(INFO)", s["issues"]["INFO"]],
        ["공급가 합계", s["supply_amt"]], ["부가세 합계", s["vat_amt"]],
        ["배송비 합계", s["shipping_fee"]], ["총 합계", s["total_amt"]],
    ]
    for status, cnt in sorted(s["order_status"].items()):
        summary_rows.append([f"주문 상태 · {STATUS_LABELS[Status(status)]}", cnt])
    _write_sheet(ws, ["항목", "값"], summary_rows)

    ow = wb.create_sheet("주문")
    orders = [order_record(o) for o in result.orders]
    _write_sheet(ow, list(orders[0].keys()), [list(r.values()) for r in orders]) if orders else None

    lw = wb.create_sheet("WMS라인")
    _write_sheet(lw, WMS_COLUMNS, [[line_record(l)[c] for c in WMS_COLUMNS] for l in result.lines])

    iw = wb.create_sheet("검수이슈")
    order = {"ERROR": 0, "WARN": 1, "INFO": 2}
    issues = sorted(result.issues, key=lambda i: (order[i.severity.value], i.row_no))
    _write_sheet(iw, ["행", "심각도", "코드", "메시지", "주문번호", "SKU"],
                 [[i.row_no, i.severity.value, i.code, i.message, i.order_id, i.sku] for i in issues])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
